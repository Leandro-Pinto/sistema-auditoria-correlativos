# -*- coding: utf-8 -*-
"""
Núcleo del Sistema de Auditoría de Correlativos – Documentos de Gestión.
Lee el Excel, interpreta N° y año del título, detecta faltantes/repetidos
y exporta el reporte a Excel. No depende de la interfaz gráfica.
"""
import json
import os
import re
from collections import defaultdict, Counter

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ESTADOS = ["CONFIRMADO", "REPETIDO", "FALTA", "SIN NÚMERO", "SIN AÑO",
           "AÑO INCOMPLETO", "AÑO INFERIDO", "FORMATO", "CORREGIDO"]
ESTADOS_REVISAR = {"SIN NÚMERO", "SIN AÑO", "AÑO INCOMPLETO",
                   "AÑO INFERIDO", "FORMATO", "CORREGIDO"}
COLORES = {  # fondo para pantalla y Excel
    "CONFIRMADO": "C6EFCE", "REPETIDO": "FFE699", "FALTA": "F8CBAD",
    "SIN NÚMERO": "F4B084", "SIN AÑO": "F4B084", "AÑO INCOMPLETO": "BDD7EE",
    "AÑO INFERIDO": "BDD7EE", "FORMATO": "E2EFDA", "CORREGIDO": "D9D2E9",
}
DESCRIPCION = {
    "CONFIRMADO": "Número y año correctos, sin repetición.",
    "REPETIDO": "El mismo tipo, número y año aparece más de una vez.",
    "FALTA": "Número ausente en la secuencia: ubicar o registrar el documento.",
    "SIN NÚMERO": "El título no tiene número: corregir.",
    "SIN AÑO": "El título no tiene año: corregir.",
    "AÑO INCOMPLETO": "Año mal escrito (ej. 202) y no se pudo deducir: corregir.",
    "AÑO INFERIDO": "Año deducido por los registros vecinos: confirmar.",
    "FORMATO": "N° y año leídos bien, pero el título tiene guiones o espacios de más.",
    "CORREGIDO": "N° o año ingresados manualmente por el auditor.",
}

RE_COMPLETO = re.compile(r"^\s*-?\s*(\d+)\s*[- ]\s*-?\s*(\d{4})(?=\s*[-/ ]|\s*$)")
RE_ANIO_CORTO = re.compile(r"^\s*-?\s*(\d+)\s*[- ]\s*-?\s*(\d{1,3})(?=\s*[-/ ])")
RE_SIN_NUMERO = re.compile(r"^\s*-\s*((?:19|20)\d{2})\s*-")
RE_SOLO_NUMERO = re.compile(r"^\s*-?\s*(\d+)\s*-")


class Registro:
    __slots__ = ("fila", "id", "tipo", "titulo", "n_leido", "y_leido",
                 "n", "y", "estado", "nota", "problema")

    def __init__(self, fila, id_, tipo, titulo):
        self.fila, self.id, self.tipo, self.titulo = fila, id_, tipo, titulo
        self.n_leido = self.y_leido = self.n = self.y = None
        self.estado = ""
        self.nota = ""
        self.problema = ""  # SIN NÚMERO / AÑO INCOMPLETO / ...


def interpretar_titulo(titulo):
    """Devuelve (numero, año, problema). problema '' si todo está bien."""
    t = (titulo or "").strip()
    m = RE_COMPLETO.match(t)
    if m:
        num, anio = int(m.group(1)), int(m.group(2))
        formato = t.startswith("-") or bool(re.match(r"^\s*-?\s*\d+\s+[- ]|^\s*-?\s*\d+\s*-\s+", t))
        return num, anio, ("FORMATO" if formato else "")
    m = RE_SIN_NUMERO.match(t)
    if m:
        return None, int(m.group(1)), "SIN NÚMERO"
    m = RE_ANIO_CORTO.match(t)
    if m:
        return int(m.group(1)), None, "AÑO INCOMPLETO"
    m = RE_SOLO_NUMERO.match(t)
    if m:
        return int(m.group(1)), None, "SIN AÑO"
    return None, None, "SIN NÚMERO"


def _norm(s):
    s = str(s or "").strip().lower()
    for a, b in (("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u")):
        s = s.replace(a, b)
    return s


def leer_excel(ruta):
    """Lee .xlsx/.csv y devuelve lista de Registro. Detecta la fila de encabezados."""
    filas = []
    if ruta.lower().endswith(".csv"):
        import csv
        with open(ruta, encoding="utf-8-sig", newline="") as f:
            muestra = f.read(4096); f.seek(0)
            delim = ";" if muestra.count(";") > muestra.count(",") else ","
            filas = [tuple(r) for r in csv.reader(f, delimiter=delim)]
    else:
        wb = load_workbook(ruta, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        filas = [tuple(r) for r in ws.iter_rows(values_only=True)]
        wb.close()

    # buscar encabezado en las primeras 20 filas
    hdr_idx, cols = None, {}
    for i, fila in enumerate(filas[:20]):
        nombres = [_norm(c) for c in fila]
        c_tit = next((j for j, c in enumerate(nombres) if c.startswith("titulo") or c == "documento"), None)
        if c_tit is not None:
            hdr_idx = i
            cols["titulo"] = c_tit
            cols["id"] = next((j for j, c in enumerate(nombres) if c in ("id", "codigo", "n", "item")), None)
            cols["tipo"] = next((j for j, c in enumerate(nombres) if c.startswith("tipo")), None)
            break
    if hdr_idx is None:
        raise ValueError("No se encontró la columna «Título» en las primeras 20 filas. "
                         "El archivo debe tener encabezados: id, Tipo de norma, Título.")

    regs = []
    for k, fila in enumerate(filas[hdr_idx + 1:], start=hdr_idx + 2):
        def val(c):
            return fila[c] if c is not None and c < len(fila) else None
        titulo = val(cols["titulo"])
        if titulo is None or str(titulo).strip() == "":
            continue
        id_ = val(cols["id"])
        try:
            id_ = int(id_)
        except (TypeError, ValueError):
            id_ = id_ if id_ not in (None, "") else f"F{k}"
        tipo = str(val(cols["tipo"]) or "Sin tipo").strip()
        regs.append(Registro(k, id_, tipo, str(titulo).strip()))
    return regs


class Auditoria:
    def __init__(self, ruta):
        self.ruta = ruta
        self.ruta_corr = os.path.splitext(ruta)[0] + ".correcciones.json"
        self.registros = leer_excel(ruta)
        self.correcciones = self._cargar_correcciones()
        self.recalcular()

    # ---------- correcciones manuales ----------
    def _cargar_correcciones(self):
        if os.path.exists(self.ruta_corr):
            try:
                with open(self.ruta_corr, encoding="utf-8") as f:
                    return json.load(f)
            except (OSError, ValueError):
                pass
        return {}

    def guardar_correccion(self, reg, n, y, nota):
        clave = str(reg.id)
        if n is None and y is None and not nota:
            self.correcciones.pop(clave, None)
        else:
            self.correcciones[clave] = {"n": n, "y": y, "nota": nota}
        with open(self.ruta_corr, "w", encoding="utf-8") as f:
            json.dump(self.correcciones, f, ensure_ascii=False, indent=1)
        self.recalcular()

    # ---------- cálculo ----------
    def recalcular(self):
        regs = self.registros
        for r in regs:
            r.n_leido, r.y_leido, r.problema = interpretar_titulo(r.titulo)
            r.n, r.y, r.nota = r.n_leido, r.y_leido, ""
        self._inferir_anios()
        for r in regs:
            c = self.correcciones.get(str(r.id))
            if c:
                if c.get("n") is not None: r.n = int(c["n"])
                if c.get("y") is not None: r.y = int(c["y"])
                r.nota = c.get("nota", "") or ""
                r.problema = "CORREGIDO"

        cuenta = Counter((r.tipo, r.y, r.n) for r in regs if r.n is not None and r.y is not None)
        for r in regs:
            if r.problema == "CORREGIDO":
                r.estado = "REPETIDO" if cuenta[(r.tipo, r.y, r.n)] > 1 else "CORREGIDO"
            elif r.y is None:
                r.estado = r.problema or "SIN AÑO"
            elif r.n is None:
                r.estado = "SIN NÚMERO"
            elif cuenta[(r.tipo, r.y, r.n)] > 1:
                r.estado = "REPETIDO"
            else:
                r.estado = r.problema or "CONFIRMADO"
        self._indexar()

    def _inferir_anios(self):
        """Para títulos sin año / año incompleto, deduce el año por los vecinos (orden de id)."""
        por_tipo = defaultdict(list)
        for r in self.registros:
            por_tipo[r.tipo].append(r)
        for tipo, lst in por_tipo.items():
            try:
                lst = sorted(lst, key=lambda r: (int(r.id), r.fila))
            except (TypeError, ValueError):
                lst = sorted(lst, key=lambda r: r.fila)
            existentes = {(r.y_leido, r.n_leido) for r in lst if r.y_leido and r.n_leido}
            for i, r in enumerate(lst):
                if r.y is not None or r.n is None:
                    continue
                prev = next((x.y_leido for x in reversed(lst[:i]) if x.y_leido), None)
                nxt = next((x.y_leido for x in lst[i + 1:] if x.y_leido), None)
                cands = [c for c in (prev, nxt) if c]
                if r.problema == "AÑO INCOMPLETO":  # ej. "202" -> años que empiezan así
                    corto = re.match(RE_ANIO_CORTO, r.titulo).group(2)
                    cands = [c for c in cands if str(c).startswith(corto)] or cands
                libres = [c for c in dict.fromkeys(cands) if (c, r.n) not in existentes]
                elegido = libres[0] if len(set(libres)) == 1 else (prev if prev == nxt else None)
                if elegido:
                    r.y = elegido
                    r.problema = "AÑO INFERIDO"

    def _indexar(self):
        self.grupos = defaultdict(lambda: defaultdict(list))
        for r in self.registros:
            if r.y is not None:
                self.grupos[(r.tipo, r.y)][r.n].append(r)

    # ---------- consultas ----------
    def tipos(self):
        return sorted({r.tipo for r in self.registros})

    def anios(self, tipo):
        return sorted({y for (t, y) in self.grupos if t == tipo})

    def secuencia(self, tipo, anio):
        """Lista de filas (n, veces, estado, títulos, ids) del 1 al último número."""
        g = self.grupos.get((tipo, anio), {})
        nums = [n for n in g if n is not None]
        maximo = max(nums) if nums else 0
        out = []
        for n in range(1, maximo + 1):
            regs = g.get(n, [])
            veces = len(regs)
            estado = "FALTA" if veces == 0 else ("REPETIDO" if veces > 1 else "CONFIRMADO")
            out.append((n, veces, estado, " | ".join(r.titulo for r in regs),
                        " | ".join(str(r.id) for r in regs)))
        for r in g.get(None, []):
            out.append(("—", 1, "SIN NÚMERO", r.titulo, str(r.id)))
        return out

    def resumen(self, tipo):
        filas = []
        for y in self.anios(tipo):
            g = self.grupos[(tipo, y)]
            nums = [n for n in g if n is not None]
            maximo = max(nums) if nums else 0
            distintos = len(nums)
            registros = sum(len(v) for v in g.values())
            rep = sum(len(v) for n, v in g.items() if n is not None and len(v) > 1)
            sin_n = len(g.get(None, []))
            falt = maximo - distintos
            comp = distintos / maximo if maximo else None
            filas.append(dict(anio=y, registros=registros, distintos=distintos, ultimo=maximo,
                              faltantes=falt, repetidos=rep, sin_numero=sin_n, completitud=comp,
                              situacion="COMPLETO" if falt == 0 else "CON FALTANTES"))
        return filas

    def faltantes(self, tipo, anio):
        return [f[0] for f in self.secuencia(tipo, anio) if f[2] == "FALTA"]

    def por_revisar(self, tipo=None):
        return [r for r in self.registros
                if (tipo is None or r.tipo == tipo) and (r.estado in ESTADOS_REVISAR or
                                                         r.problema in ESTADOS_REVISAR)]

    def sin_anio(self, tipo):
        return [r for r in self.registros if r.tipo == tipo and r.y is None]

    # ---------- exportar ----------
    def exportar_excel(self, destino, tipos=None, progreso=None):
        tipos = tipos or self.tipos()
        wb = Workbook()
        NAVY = "1F4E78"
        f_hdr = Font(name="Arial", bold=True, color="FFFFFF")
        fill_hdr = PatternFill("solid", fgColor=NAVY)
        f_body = Font(name="Arial", size=10)
        fills = {k: PatternFill("solid", fgColor=v) for k, v in COLORES.items()}

        def encabezado(ws, cols, anchos):
            ws.append(cols)
            for c in ws[ws.max_row]:
                c.font, c.fill = f_hdr, fill_hdr
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for i, w in enumerate(anchos, 1):
                ws.column_dimensions[get_column_letter(i)].width = w
            ws.freeze_panes = f"A{ws.max_row + 1}"

        filas_ws = {}

        def fila(ws, valores, estado=None, pintar_todo=True, col_estado=None):
            r = filas_ws.get(ws.title) or ws.max_row
            r += 1
            filas_ws[ws.title] = r
            celdas = []
            for j, v in enumerate(valores, 1):
                c = ws.cell(r, j, v)
                c.font = f_body
                celdas.append(c)
            if estado in fills and estado != "CONFIRMADO":
                objetivo = celdas if pintar_todo else [celdas[col_estado]]
                for c in objetivo:
                    c.fill = fills[estado]

        # Resumen
        ws = wb.active; ws.title = "Resumen"
        ws["A1"] = "Sistema de Auditoría de Correlativos – Documentos de Gestión"
        ws["A1"].font = Font(name="Arial", bold=True, size=14, color=NAVY)
        ws["A2"] = f"Archivo auditado: {os.path.basename(self.ruta)}"
        ws["A2"].font = Font(name="Arial", italic=True, size=9)
        ws.append([])
        encabezado(ws, ["Tipo de documento", "Año", "Registros", "N° distintos", "Último N°",
                        "Faltantes", "Registros repetidos", "Sin número", "% Completitud", "Situación"],
                   [32, 8, 11, 12, 11, 11, 13, 11, 13, 16])
        ws.freeze_panes = "A5"
        primera = ws.max_row + 1
        for t in tipos:
            for s in self.resumen(t):
                fila(ws, [t, s["anio"], s["registros"], s["distintos"], s["ultimo"], s["faltantes"],
                          s["repetidos"], s["sin_numero"], s["completitud"], s["situacion"]],
                     "FALTA" if s["faltantes"] else "CONFIRMADO", pintar_todo=False, col_estado=9)
                ws.cell(ws.max_row, 9).number_format = "0.0%"
        ultima = ws.max_row
        ws.append(["TOTAL", None] + [f"=SUM({get_column_letter(c)}{primera}:{get_column_letter(c)}{ultima})"
                                     for c in range(3, 9)] +
                  [f"=IF(SUM(E{primera}:E{ultima})=0,\"\",D{ultima+1}/SUM(E{primera}:E{ultima}))", None])
        ws.cell(ultima + 1, 5).value = None
        for c in ws[ws.max_row]:
            c.font = Font(name="Arial", bold=True)
        ws.cell(ws.max_row, 9).number_format = "0.0%"
        ws.append([])
        ws.append(["Leyenda de estados"]); ws.cell(ws.max_row, 1).font = Font(name="Arial", bold=True)
        for e in ESTADOS:
            ws.append([e, DESCRIPCION[e]])
            ws.cell(ws.max_row, 1).fill = fills[e]
            for c in ws[ws.max_row]: c.font = f_body

        # Faltantes
        wf = wb.create_sheet("Faltantes")
        encabezado(wf, ["Tipo de documento", "Año", "N° faltante"], [32, 8, 12])
        for t in tipos:
            for y in self.anios(t):
                for n in self.faltantes(t, y):
                    fila(wf, [t, y, n], "FALTA")
        wf.auto_filter.ref = f"A1:C{max(wf.max_row, 2)}"

        # Por revisar
        wr = wb.create_sheet("Por revisar")
        encabezado(wr, ["Fila origen", "id", "Tipo de documento", "Título", "N° final", "Año final",
                        "Estado", "Detalle", "Nota del auditor"], [10, 11, 30, 32, 9, 9, 17, 55, 40])
        for t in tipos:
            for r in self.por_revisar(t):
                det = r.problema if r.problema and r.problema != r.estado else r.estado
                fila(wr, [r.fila, r.id, r.tipo, r.titulo, r.n, r.y, r.estado,
                          DESCRIPCION.get(det, ""), r.nota], r.problema or r.estado)
        wr.auto_filter.ref = f"A1:I{max(wr.max_row, 2)}"

        # Todos + hojas por año
        cols = ["id", "Tipo de documento", "Título", "N° documento", "Año", "Estado", "Nota del auditor"]
        anchos = [11, 30, 32, 12, 8, 17, 45]
        wt = wb.create_sheet("Todos")
        encabezado(wt, cols, anchos)
        total = sum(len(self.anios(t)) for t in tipos) or 1
        hechos = 0
        for t in tipos:
            sigla = "".join(p[0] for p in t.split() if p[0].isalpha()).upper()[:6] or "DOC"
            for y in self.anios(t):
                wy = wb.create_sheet(f"{sigla} {y}"[:31])
                encabezado(wy, cols, anchos)
                g = self.grupos[(t, y)]
                nums = [n for n in g if n is not None]
                for n in range(1, (max(nums) if nums else 0) + 1):
                    regs = g.get(n)
                    if not regs:
                        vals, est = [None, t, "", n, y, "FALTA", ""], "FALTA"
                        fila(wt, vals, est); fila(wy, vals, est)
                    for r in regs or []:
                        est = r.estado if r.estado in ("REPETIDO", "FALTA") else (r.problema or r.estado)
                        vals = [r.id, r.tipo, r.titulo, r.n, r.y, r.estado, r.nota]
                        fila(wt, vals, est); fila(wy, vals, est)
                for r in g.get(None, []):
                    vals = [r.id, r.tipo, r.titulo, None, r.y, r.estado, r.nota]
                    fila(wt, vals, "SIN NÚMERO"); fila(wy, vals, "SIN NÚMERO")
                wy.auto_filter.ref = f"A1:G{max(wy.max_row, 2)}"
                hechos += 1
                if progreso:
                    progreso(hechos / total)
            for r in self.sin_anio(t):
                fila(wt, [r.id, r.tipo, r.titulo, r.n, None, r.estado, r.nota], r.estado)
        wt.auto_filter.ref = f"A1:G{max(wt.max_row, 2)}"
        wb.save(destino)
        return destino
